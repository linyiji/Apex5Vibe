from copy import deepcopy
from io import BytesIO

import pytest
from apexdefend.config import HANDOFF
from apexdefend.db import Reference
from apexdefend.errors import AppError
from apexdefend.field_sheets import field_sheet
from apexdefend.parsers import parse_evidence, parse_invoice
from apexdefend.references import select_reference
from conftest import complete, key, new_case, upload
from openpyxl import load_workbook
from sqlalchemy import select


def case2(system):
    client, worker, analyzer, _ = system
    analyzer.rules["relationship"] = "TRADE_SCALE_HIGH"
    case_id = new_case(client)
    upload(client, case_id, worker, 2)
    return case_id, complete(client, worker, case_id)


def sheet(client, case_id, source_name, old=None, data=None, filename="same-name.xlsx", headers=None):
    url = f"/v1/cases/{case_id}/modules/invoice_integrity"
    module = client.get(url).json()
    values = {
        "expected_module_version": module["module_version"],
        "rerun": "true",
        "evidence_group": "invoice_field_sheet",
    }
    if old:
        values["supersedes_evidence_id"] = old
    response = client.post(
        url + "/evidence",
        headers=headers or key(),
        files=[("files", (filename, data or (HANDOFF / f"source/{source_name}").read_bytes()))],
        data=values,
    )
    return response


def score(ws):
    return ws["aggregate"]["current_result"]["final_risk_score"]


def test_case2_errors_then_explicit_correction_isolated_and_auditable(system):
    client, worker, analyzer, _ = system
    case_id, initial = case2(system)
    url = f"/v1/cases/{case_id}/modules/invoice_integrity"
    assert initial["invoice"]["total"] == "50000.00" and score(initial) == 35
    analyzer.rules["invoice_integrity"] = "INTEGRITY_CONFLICT"
    response = sheet(client, case_id, "ApexDefend_Case_2_Errors.xlsx", filename="Corrected.xlsx")
    assert response.status_code == 202, response.text
    bad_id = response.json()["evidence_ids"][0]
    assert client.get(f"/v1/cases/{case_id}").json()["aggregate"]["decision"]["status"] == "pending"
    assert worker.run_once() and worker.run_once()
    errors = client.get(f"/v1/cases/{case_id}").json()
    assert score(errors) == 50 and errors["invoice"]["total"] == "50000.00"
    evidence = analyzer.calls[-1]["evidence"][0]
    assert evidence["source_id"] == bad_id
    mapped = evidence["content"]["field_sheet"]
    assert {d["cell"] for d in mapped["differences"]} == {"C25", "C51"}
    assert mapped["fields"]["invoice_total"]["normalized_value"] == "80000.00"
    assert next(c for c in mapped["checks"] if c["code"] == "subtotal_plus_charges")["passed"] is False
    before = client.get(url).json()
    duplicate = sheet(client, case_id, "ApexDefend_Case_2_Errors.xlsx", filename="renamed.xlsx")
    assert duplicate.json()["evidence_ids"] == [bad_id]
    assert client.get(url).json()["input_version"] == before["input_version"]
    assert worker.run_once()  # Explicit rerun of the same current evidence, no reparse.
    assert sheet(client, case_id, "ApexDefend_Case_2_Corrected.xlsx").status_code == 409
    corrected = sheet(
        client, case_id, "ApexDefend_Case_2_Corrected.xlsx", old=bad_id, filename="Corrected.xlsx"
    )
    assert corrected.status_code == 202, corrected.text
    good_id = corrected.json()["evidence_ids"][0]
    pending = client.get(url).json()["evidence"]
    assert next(e for e in pending if e["evidence_id"] == bad_id)["status"] == "active"
    assert next(e for e in pending if e["evidence_id"] == good_id)["status"] == "pending"
    analyzer.rules["invoice_integrity"] = "INTEGRITY_OK"
    assert worker.run_once() and worker.run_once()
    final = client.get(f"/v1/cases/{case_id}").json()
    assert score(final) == 35 and final["aggregate"]["decision"]["status"] == "manual_review"
    assert [e["source_id"] for e in analyzer.calls[-1]["evidence"]] == [good_id]
    assert analyzer.calls[-1]["evidence"][0]["content"]["field_sheet"]["differences"] == []
    history = client.get(url).json()["evidence"]
    old = next(e for e in history if e["evidence_id"] == bad_id)
    new = next(e for e in history if e["evidence_id"] == good_id)
    assert old["status"] == "superseded" and new["status"] == "active"
    assert new["version"] == 2 and new["supersedes_evidence_id"] == bad_id
    assert client.get(f"/v1/cases/{case_id}/files/{bad_id}").status_code == 200
    for first, last in zip(initial["modules"], final["modules"]):
        if first["module_key"] != "invoice_integrity":
            assert first["current_run"]["run_id"] == last["current_run"]["run_id"]


def invalid_sheet():
    book = load_workbook(HANDOFF / "source/ApexDefend_Case_2_Corrected.xlsx")
    book["Invoice Fields"]["C38"] = "not a valid price"
    result = BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


def test_failed_replacement_retains_old_evidence_and_can_be_recovered(system):
    client, worker, analyzer, service = system
    case_id, _ = case2(system)
    analyzer.rules["invoice_integrity"] = "INTEGRITY_CONFLICT"
    first = sheet(client, case_id, "ApexDefend_Case_2_Errors.xlsx").json()["evidence_ids"][0]
    assert worker.run_once() and worker.run_once()
    count = len(analyzer.calls)
    bad = sheet(client, case_id, "", old=first, data=invalid_sheet()).json()["evidence_ids"][0]
    assert worker.run_once() and not worker.run_once()
    assert len(analyzer.calls) == count
    ws = client.get(f"/v1/cases/{case_id}").json()
    assert ws["aggregate"]["current_result"] is None
    assert ws["aggregate"]["last_valid_result"]["final_risk_score"] == 50
    module = next(m for m in ws["modules"] if m["module_key"] == "invoice_integrity")
    assert module["state"] == "needs_info"
    assert next(e for e in module["evidence"] if e["evidence_id"] == first)["status"] == "active"
    assert next(e for e in module["evidence"] if e["evidence_id"] == bad)["status"] == "failed"
    analyzer.rules["invoice_integrity"] = "INTEGRITY_OK"
    assert sheet(client, case_id, "ApexDefend_Case_2_Corrected.xlsx", old=first).status_code == 202
    assert worker.run_once() and worker.run_once()
    assert score(client.get(f"/v1/cases/{case_id}").json()) == 35


def test_replacement_rejects_cross_case_and_pending_versions(system):
    client, worker, _, _ = system
    first_case, _ = case2(system)
    accepted = sheet(client, first_case, "ApexDefend_Case_2_Errors.xlsx")
    evidence_id = accepted.json()["evidence_ids"][0]
    assert sheet(client, first_case, "ApexDefend_Case_2_Corrected.xlsx").status_code == 409
    assert worker.run_once() and worker.run_once()
    second_case, _ = case2(system)
    assert sheet(client, second_case, "ApexDefend_Case_2_Corrected.xlsx", old=evidence_id).status_code == 409


def test_field_mapping_follows_keys_not_rows_and_rejects_duplicates():
    invoice = parse_invoice((HANDOFF / "source/ApexDefend_Case_2_Invoice.pdf").read_bytes(), "pdf")
    content = parse_evidence("anything.xlsx", (HANDOFF / "source/ApexDefend_Case_2_Errors.xlsx").read_bytes())
    for cell in content["cells"]:
        for column in ["A", "B", "C", "D"]:
            if cell["cell"] == column + "25":
                cell["cell"] = column + "80"
    mapped = field_sheet(content, "sheet", invoice)
    assert mapped["fields"]["buyer_address"]["cell"] == "C80"
    assert mapped["fields"]["parties_related"]["normalized_value"] is False
    content["cells"].append({"sheet": "Invoice Fields", "cell": "B90", "value": "invoice_total"})
    with pytest.raises(AppError) as error:
        field_sheet(content, "sheet", invoice)
    assert error.value.code == "INVALID_FIELD_SHEET"


def test_reference_selection_excludes_unrelated_future_and_current_records(system):
    _, _, _, service = system
    _, ws = case2(system)
    invoice = ws["invoice"]
    with service.db.sessions() as session:
        trade = deepcopy(session.get(Reference, "ref_trade_history_v1").document)
        base = trade["data"]["trades"][0]
        for change in [
            {"buyer_id": "other"},
            {"seller_id": "other"},
            {"currency": "USD"},
            {"sku": "other"},
            {"status": "pending"},
            {"invoice_date": "2026-12-01"},
            {"invoice_number": invoice["invoice_number"]},
        ]:
            trade["data"]["trades"].append({**base, **change})
        selected = select_reference(trade, "relationship", invoice, "2026-10-04")
        assert len(selected["data"]["trades"]) == 12
        for reference in session.scalars(select(Reference).where(Reference.kind == "market_benchmark")):
            wrong = deepcopy(reference.document)
            wrong["data"]["currency"] = "USD"
            assert select_reference(wrong, "market_price", invoice, "2026-10-04") is None
        payments = deepcopy(session.get(Reference, "ref_payments_v1").document)
        payments["data"]["records"].append({**payments["data"]["records"][0], "buyer_id": "other"})
        assert (
            len(select_reference(payments, "payment_history", invoice, "2026-10-04")["data"]["records"]) == 12
        )


def test_missing_comparable_reference_stops_before_ai(system):
    client, worker, analyzer, service = system
    case_id = new_case(client)
    upload(client, case_id, worker, 2)
    with service.db.transaction() as session:
        record = session.get(Reference, "ref_market_v1")
        doc = deepcopy(record.document)
        doc["data"]["sku"] = "unrelated"
        record.document = doc
    ws = complete(client, worker, case_id)
    assert not any(i["module_key"] == "market_price" for i in analyzer.calls)
    module = next(m for m in ws["modules"] if m["module_key"] == "market_price")
    assert module["state"] == "needs_info" and module["final_score"] is None


def test_ai_failure_after_activation_does_not_claim_correction_success(system):
    client, worker, analyzer, _ = system
    case_id, _ = case2(system)
    analyzer.rules["invoice_integrity"] = "INTEGRITY_CONFLICT"
    original = sheet(client, case_id, "ApexDefend_Case_2_Errors.xlsx").json()["evidence_ids"][0]
    assert worker.run_once() and worker.run_once()
    replacement = sheet(client, case_id, "ApexDefend_Case_2_Corrected.xlsx", old=original).json()[
        "evidence_ids"
    ][0]
    analyzer.fail_key = "invoice_integrity"
    assert worker.run_once() and worker.run_once()
    ws = client.get(f"/v1/cases/{case_id}").json()
    module = next(m for m in ws["modules"] if m["module_key"] == "invoice_integrity")
    assert module["state"] == "failed" and ws["aggregate"]["current_result"] is None
    assert ws["aggregate"]["decision"]["status"] == "pending"
    assert ws["aggregate"]["last_valid_result"]["final_risk_score"] == 50
    assert next(e for e in module["evidence"] if e["evidence_id"] == replacement)["status"] == "active"
    assert next(e for e in module["evidence"] if e["evidence_id"] == original)["status"] == "superseded"


def test_financing_selection_uses_identity_decimal_amount_and_active_status(system):
    _, ws = case2(system)
    invoice = ws["invoice"]
    match = {
        "invoice_number": invoice["invoice_number"],
        "seller_tax_id": invoice["seller"]["tax_id"],
        "buyer_tax_id": invoice["buyer"]["tax_id"],
        "currency": "GBP",
        "total": "50000",
        "invoice_date": invoice["invoice_date"],
        "status": "active",
        "financed_at": "2026-10-03",
    }
    doc = {
        "source_id": "registry",
        "kind": "financing_registry",
        "as_of": "2026-10-04",
        "data": {
            "coverage_complete_for_demo": True,
            "records": [
                match,
                {**match, "status": "settled"},
                {**match, "buyer_tax_id": "other"},
                {**match, "financed_at": "2026-10-05"},
            ],
        },
    }
    selected = select_reference(doc, "duplicate_financing", invoice, "2026-10-04")
    assert selected["data"]["records"] == [match]
